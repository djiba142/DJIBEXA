import csv
import io
import re
import unicodedata
import zipfile
from datetime import date, datetime
from decimal import Decimal, InvalidOperation

from django.db.models import Sum
from openpyxl import load_workbook
from rest_framework.exceptions import ValidationError
from xlrd.biffh import XLRDError

from .models import PurchaseReceiptItem

MAX_FILE_SIZE = 10 * 1024 * 1024
MAX_IMPORT_ROWS = 1000

FIELD_ALIASES = {
    'product_reference': (
        'reference article nexora', 'reference nexora', 'reference interne',
        'product reference', 'product code', 'sku', 'code produit', 'reference',
    ),
    'supplier_reference': (
        'reference article fournisseur', 'reference fournisseur',
        'supplier reference', 'supplier sku', 'article fournisseur',
    ),
    'designation': (
        'designation description', 'designation', 'description', 'product name',
        'name', 'produit', 'article',
    ),
    'quantity_ordered': ('quantite commandee', 'ordered quantity', 'ordered qty'),
    'quantity_received': (
        'quantite reellement recue', 'quantite recue', 'received quantity',
        'received qty', 'quantity received', 'qty received', 'qte recue',
        'quantite', 'quantity', 'qty',
    ),
    'quantity_accepted': ('quantite acceptee', 'accepted quantity', 'qty accepted'),
    'quantity_quarantined': (
        'quantite quarantaine', 'quantite a isoler', 'a isoler', 'quantity quarantined',
        'quarantined quantity',
    ),
    'unit_cost': (
        'prix achat unitaire', 'prix d achat unitaire', 'unit purchase price',
        'purchase unit price', 'unit cost', 'purchase price',
    ),
    'supplier_barcode': (
        'code barre fournisseur', 'code barres fournisseur', 'supplier barcode',
        'barcode',
    ),
    'items_per_carton': ('pieces par carton', 'items per carton', 'units per carton'),
    'discrepancy_reason': (
        'motif ecart reserves', 'motif ecart', 'discrepancy reason',
        'reason for discrepancy',
    ),
    'condition_status': ('etat conformite', 'condition status', 'conformity'),
    'packaging_condition': (
        'etat emballage qualite', 'packaging condition', 'packaging quality',
    ),
    'supplier_lot_reference': (
        'no lot ref livraison', 'n lot ref livraison', 'numero lot', 'supplier lot reference',
        'delivery reference',
    ),
    'line_notes': ('commentaire', 'comment', 'notes', 'observation'),
    'location': ('emplacement destination', 'destination', 'location', 'warehouse'),
}

METADATA_ALIASES = {
    'source_reference': ('no bon de reception interne', 'receipt reference'),
    'received_date': ('date de reception', 'received date', 'receipt date'),
    'supplier_name': ('fournisseur', 'supplier'),
    'purchase_order_reference': (
        'no commande fournisseur', 'commande fournisseur', 'purchase order',
        'purchase order reference',
    ),
    'supplier_document_reference': (
        'facture bon de livraison fournisseur', 'reference document fournisseur',
        'supplier invoice', 'supplier delivery note',
    ),
    'currency': ('devise achat', 'purchase currency', 'currency'),
    'location': ('site de destination', 'destination', 'storage location'),
}


def _normalize(value):
    normalized = unicodedata.normalize('NFKD', str(value or '').strip().lower())
    normalized = ''.join(char for char in normalized if not unicodedata.combining(char))
    return re.sub(r'[^a-z0-9]+', ' ', normalized).strip()


NORMALIZED_FIELD_ALIASES = {
    field: {_normalize(alias) for alias in aliases}
    for field, aliases in FIELD_ALIASES.items()
}
NORMALIZED_METADATA_ALIASES = {
    field: {_normalize(alias) for alias in aliases}
    for field, aliases in METADATA_ALIASES.items()
}


def _clean(value):
    if value is None:
        return ''
    if isinstance(value, (datetime, date)):
        return value.isoformat()
    return str(value).strip()


def _normalize_received_date(value):
    raw = _clean(value)
    if not raw:
        return ''
    try:
        return datetime.fromisoformat(raw).isoformat()
    except ValueError:
        for date_format in ('%d/%m/%Y', '%d-%m-%Y', '%Y-%m-%d'):
            try:
                parsed_date = datetime.strptime(raw, date_format).date()
                return datetime.combine(parsed_date, datetime.min.time()).isoformat()
            except ValueError:
                continue
    return ''


def _read_rows(uploaded_file):
    uploaded_file.seek(0)
    content = uploaded_file.read()
    if len(content) > MAX_FILE_SIZE:
        raise ValidationError({'file': 'Le fichier ne doit pas dépasser 10 Mo.'})
    filename = str(getattr(uploaded_file, 'name', '')).lower()
    try:
        if filename.endswith('.csv'):
            text = content.decode('utf-8-sig')
            try:
                dialect = csv.Sniffer().sniff(text[:4096], delimiters=',;\t')
            except csv.Error:
                dialect = csv.excel
            reader = csv.reader(io.StringIO(text), dialect=dialect)
            return [
                [value for value in row]
                for row in reader
            ]
        if not filename.endswith('.xlsx'):
            raise ValidationError({'file': 'Utilisez un fichier Excel .xlsx ou CSV.'})
        workbook = load_workbook(io.BytesIO(content), read_only=True, data_only=True)
    except (UnicodeDecodeError, OSError, ValueError, zipfile.BadZipFile, XLRDError) as exc:
        raise ValidationError({'file': f'Le fichier ne peut pas être analysé : {exc}'}) from exc

    try:
        return [
            (sheet.title, [list(row) for row in sheet.iter_rows(values_only=True)])
            for sheet in workbook.worksheets
        ]
    finally:
        workbook.close()


def _find_table(sheets):
    for sheet_name, rows in sheets:
        for row_index, row in enumerate(rows):
            normalized_headers = [_normalize(value) for value in row]
            mapping = {}
            for index, header in enumerate(normalized_headers):
                for field, aliases in NORMALIZED_FIELD_ALIASES.items():
                    if header in aliases:
                        mapping.setdefault(field, []).append(index)
            has_identifier = any(
                field in mapping
                for field in ('product_reference', 'supplier_reference', 'designation')
            )
            if has_identifier and 'quantity_received' in mapping:
                if any(len(indices) > 1 for indices in mapping.values()):
                    raise ValidationError({
                        'file': 'Le tableau contient plusieurs colonnes correspondant au même champ.',
                    })
                return sheet_name, rows, row_index, {
                    field: indices[0] for field, indices in mapping.items()
                }
    raise ValidationError({
        'file': (
            'Aucun tableau de réception reconnu. Le fichier doit contenir une référence '
            'ou une désignation et une colonne de quantité réellement reçue.'
        ),
    })


def _read_metadata(rows, header_row_index):
    metadata = {}
    label_fields = {
        alias: field
        for field, aliases in NORMALIZED_METADATA_ALIASES.items()
        for alias in aliases
    }
    for row in rows[:header_row_index]:
        labels = [
            (index, label_fields[_normalize(value)])
            for index, value in enumerate(row)
            if _normalize(value) in label_fields
        ]
        for label_index, field in labels:
            next_label = next(
                (index for index, _ in labels if index > label_index),
                len(row),
            )
            value = next(
                (
                    _clean(row[index])
                    for index in range(label_index + 1, next_label)
                    if _clean(row[index])
                ),
                '',
            )
            if value:
                metadata[field] = value
    return metadata


def _parse_integer(value, label, line_number, errors, allow_zero=True):
    raw = _clean(value).replace(',', '.')
    if not raw:
        errors.append(f'Ligne {line_number} : {label} manquante.')
        return None
    try:
        parsed = Decimal(raw)
        if not parsed.is_finite() or parsed != parsed.to_integral_value():
            raise InvalidOperation
        result = int(parsed)
        if result < (0 if allow_zero else 1):
            raise InvalidOperation
        return result
    except (InvalidOperation, ValueError):
        errors.append(f'Ligne {line_number} : {label} doit être un entier valide.')
        return None


def _parse_decimal(value, label, line_number, errors):
    raw = _clean(value)
    if not raw:
        return None
    raw = re.sub(r'\s|€|\$|GNF|USD', '', raw, flags=re.IGNORECASE)
    if ',' in raw and '.' in raw:
        decimal_separator = ',' if raw.rfind(',') > raw.rfind('.') else '.'
        grouping_separator = '.' if decimal_separator == ',' else ','
        raw = raw.replace(grouping_separator, '').replace(decimal_separator, '.')
    else:
        raw = raw.replace(',', '.')
    try:
        value = Decimal(raw)
        if not value.is_finite() or value < 0:
            raise InvalidOperation
        return str(value)
    except (InvalidOperation, ValueError):
        errors.append(f'Ligne {line_number} : {label} doit être un montant positif ou nul.')
        return None


def _product_identifiers(product):
    identifiers = {product.sku}
    if product.manufacturer_reference:
        identifiers.add(product.manufacturer_reference)
    identifiers.update(
        value.strip()
        for value in re.split(r'[,;\n|]+', product.alternative_references or '')
        if value.strip()
    )
    return {_normalize(identifier) for identifier in identifiers if identifier}


def _match_product(raw, order_items):
    references = [
        _normalize(raw.get('product_reference')),
        _normalize(raw.get('supplier_reference')),
    ]
    references = [reference for reference in references if reference]
    if references:
        matching = [
            item for item in order_items
            if _product_identifiers(item.product).intersection(references)
        ]
        if len(matching) == 1:
            return matching[0].product, ''
        if len(matching) > 1:
            return None, 'Référence ambiguë parmi les produits de la commande.'

    designation = _normalize(raw.get('designation'))
    if designation:
        matching = [
            item for item in order_items
            if _normalize(item.product.name) == designation
        ]
        if len(matching) == 1:
            return matching[0].product, ''
        if len(matching) > 1:
            return None, 'Désignation ambiguë parmi les produits de la commande.'
    return None, 'Produit non reconnu dans la commande ; aucun produit ne sera créé automatiquement.'


def parse_receipt_import(uploaded_file, order, default_location, locations):
    sheets = _read_rows(uploaded_file)
    if sheets and isinstance(sheets[0], tuple):
        table = _find_table(sheets)
    else:
        table = _find_table([('CSV', sheets)])
    _, rows, header_row_index, mapping = table
    if sum(
        any(_clean(value) for value in row)
        for row in rows[header_row_index + 1:]
    ) > MAX_IMPORT_ROWS:
        raise ValidationError({
            'file': f'Le fichier ne peut pas contenir plus de {MAX_IMPORT_ROWS} lignes de réception.',
        })
    metadata = _read_metadata(rows, header_row_index)

    errors = []
    raw_received_date = metadata.get('received_date', '')
    received_date = _normalize_received_date(raw_received_date)
    if raw_received_date and not received_date:
        errors.append('La date de réception dans le fichier est invalide.')
    if metadata.get('supplier_name') and _normalize(metadata['supplier_name']) != _normalize(order.supplier.name):
        errors.append('Le fournisseur du fichier ne correspond pas à la commande choisie.')
    if (
        metadata.get('purchase_order_reference')
        and _normalize(metadata['purchase_order_reference']) != _normalize(order.reference)
    ):
        errors.append('Le numéro de commande du fichier ne correspond pas à la commande choisie.')
    if metadata.get('currency') and _normalize(metadata['currency']) != _normalize(order.currency):
        errors.append('La devise du fichier ne correspond pas à la devise de la commande.')

    locations_by_identifier = {}
    ambiguous_locations = set()
    for location in locations:
        for identifier in (location.name, location.code):
            key = _normalize(identifier)
            if key in locations_by_identifier and locations_by_identifier[key].pk != location.pk:
                ambiguous_locations.add(key)
            else:
                locations_by_identifier[key] = location
    if default_location is None and metadata.get('location'):
        metadata_location = _normalize(metadata['location'])
        if metadata_location not in ambiguous_locations:
            default_location = locations_by_identifier.get(metadata_location)
        if default_location is None:
            errors.append('L’emplacement indiqué dans le fichier est absent ou non autorisé.')

    order_items = list(order.items.select_related('product').all())
    previously_received = {
        item['product_id']: item['total']
        for item in PurchaseReceiptItem.objects.filter(
            purchase_receipt__purchase_order=order,
        ).exclude(
            purchase_receipt__status='cancelled',
        ).values('product_id').annotate(total=Sum('quantity_received'))
    }
    imported_quantities = {}
    seen_products = set()
    preview_rows = []
    for row_index, row in enumerate(rows[header_row_index + 1:], start=header_row_index + 2):
        values = {
            field: row[index] if index < len(row) else None
            for field, index in mapping.items()
        }
        if not any(_clean(value) for value in values.values()):
            continue
        row_errors = []
        product, product_error = _match_product(values, order_items)
        if product_error:
            row_errors.append(f'Ligne {row_index} : {product_error}')
        elif product.pk in seen_products:
            row_errors.append(f'Ligne {row_index} : produit en double dans le fichier.')
        else:
            seen_products.add(product.pk)

        quantity = _parse_integer(
            values.get('quantity_received'),
            'la quantité réellement reçue',
            row_index,
            row_errors,
        )
        if product is not None and quantity is not None:
            imported_quantities[product.pk] = imported_quantities.get(product.pk, 0) + quantity
            order_line = next(item for item in order_items if item.product_id == product.pk)
            remaining = max(
                order_line.quantity - previously_received.get(product.pk, 0),
                0,
            )
            if imported_quantities[product.pk] > remaining:
                row_errors.append(
                    f'Ligne {row_index} : la quantité reçue dépasse le reliquat de {remaining} sur la commande.'
                )
        ordered_quantity = None
        if _clean(values.get('quantity_ordered')):
            ordered_quantity = _parse_integer(
                values.get('quantity_ordered'),
                'la quantité commandée',
                row_index,
                row_errors,
            )
            if product is not None and ordered_quantity is not None:
                order_line = next(item for item in order_items if item.product_id == product.pk)
                if ordered_quantity != order_line.quantity:
                    row_errors.append(
                        f'Ligne {row_index} : quantité commandée différente de la commande DJIBEXA.'
                    )

        quarantined = 0
        if _clean(values.get('quantity_quarantined')):
            quarantined = _parse_integer(
                values.get('quantity_quarantined'),
                'la quantité à isoler',
                row_index,
                row_errors,
            )
        accepted = quantity
        if _clean(values.get('quantity_accepted')):
            accepted = _parse_integer(
                values.get('quantity_accepted'),
                'la quantité acceptée',
                row_index,
                row_errors,
            )
        elif quantity is not None:
            accepted = quantity - (quarantined or 0)
        if (
            quantity is not None
            and accepted is not None
            and accepted + (quarantined or 0) != quantity
        ):
            row_errors.append(
                f'Ligne {row_index} : quantité acceptée et isolée différente de la quantité reçue.'
            )

        unit_cost = _parse_decimal(values.get('unit_cost'), "le prix d'achat unitaire", row_index, row_errors)
        carton_size = 1
        if _clean(values.get('items_per_carton')):
            carton_size = _parse_integer(
                values.get('items_per_carton'),
                'le nombre de pièces par carton',
                row_index,
                row_errors,
                allow_zero=False,
            ) or 1

        location = default_location
        raw_location = _clean(values.get('location'))
        if raw_location:
            location_key = _normalize(raw_location)
            location = locations_by_identifier.get(location_key)
            if location_key in ambiguous_locations:
                row_errors.append(f'Ligne {row_index} : emplacement ambigu, utilisez son code.')
            elif location is None:
                row_errors.append(
                    f'Ligne {row_index} : emplacement absent, inactif ou non autorisé.'
                )
        if location is None:
            row_errors.append(f'Ligne {row_index} : choisissez un emplacement de réception.')

        preview_rows.append({
            'line': row_index,
            'product': product.pk if product else None,
            'product_reference': product.sku if product else _clean(values.get('product_reference')),
            'product_name': product.name if product else _clean(values.get('designation')),
            'quantity_ordered': ordered_quantity,
            'quantity_received': quantity,
            'quantity_accepted': accepted,
            'quantity_quarantined': quarantined,
            'unit_cost': unit_cost,
            'supplier_reference': _clean(values.get('supplier_reference')),
            'supplier_barcode': _clean(values.get('supplier_barcode')),
            'items_per_carton': carton_size,
            'supplier_lot_reference': _clean(values.get('supplier_lot_reference')),
            'condition_status': _clean(values.get('condition_status')),
            'packaging_condition': _clean(values.get('packaging_condition')),
            'discrepancy_reason': _clean(values.get('discrepancy_reason')),
            'notes': _clean(values.get('line_notes')),
            'location': location.pk if location else None,
            'location_name': location.name if location else raw_location,
            'errors': row_errors,
        })

    if not preview_rows:
        errors.append(
            'Aucune ligne de produit exploitable. Vérifiez les quantités réellement reçues dans le fichier.'
        )
    if any(row['errors'] for row in preview_rows):
        errors.append('Corrigez les lignes signalées avant de créer la réception.')

    return {
        'source_name': str(getattr(uploaded_file, 'name', 'reception.xlsx')),
        'source_reference': metadata.get('source_reference', ''),
        'supplier_name': order.supplier.name,
        'purchase_order': order.pk,
        'purchase_order_reference': order.reference,
        'supplier_document_reference': metadata.get('supplier_document_reference', ''),
        'received_date': received_date,
        'currency': metadata.get('currency', order.currency),
        'rows': preview_rows,
        'total_rows': len(preview_rows),
        'valid_rows': sum(not row['errors'] for row in preview_rows),
        'errors': errors,
    }
